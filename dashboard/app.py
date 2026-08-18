import calendar
from datetime import datetime, timedelta
import io
from pathlib import Path
import sqlite3

from flask import Flask, render_template, request, Response
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)

BASE_DIR = Path(__file__).resolve().parent
DB_FILE = BASE_DIR.parent / "attendance.db"


def get_db():
    db = sqlite3.connect(DB_FILE)
    db.row_factory = sqlite3.Row
    return db


def compute_total_hours(min_ts, max_ts, total_punches):
    if total_punches > 1 and min_ts and max_ts:
        try:
            t_in = datetime.fromisoformat(min_ts)
            t_out = datetime.fromisoformat(max_ts)
            diff = t_out - t_in
            secs = int(diff.total_seconds())
            if secs > 0:
                h = secs // 3600
                m = (secs % 3600) // 60
                return f"{h}h {m:02d}m"
        except Exception:
            pass
    return "N/A"


@app.route("/")
def attendance():
    selected_user_id = request.args.get("user_id", "").strip()
    selected_from_date = request.args.get("from_date", "").strip()
    selected_to_date = request.args.get("to_date", "").strip()

    try:
        page = max(1, int(request.args.get("page", 1)))
    except ValueError:
        page = 1

    per_page = 25
    offset = (page - 1) * per_page

    db = get_db()

    # Fetch all registered users
    all_users = db.execute(
        "SELECT user_id, name FROM users ORDER BY CAST(user_id AS INTEGER), user_id"
    ).fetchall()

    # Determine KPI analytics target date (selected date or latest date in DB)
    if selected_from_date and selected_to_date and selected_from_date == selected_to_date:
        target_date = selected_from_date
    elif selected_from_date and not selected_to_date:
        target_date = selected_from_date
    else:
        latest_date_row = db.execute("SELECT MAX(DATE(timestamp)) FROM attendance").fetchone()
        target_date = latest_date_row[0] if latest_date_row and latest_date_row[0] else datetime.now().strftime("%Y-%m-%d")

    # Compute KPI analytics for target date
    present_rows = db.execute(
        "SELECT DISTINCT user_id FROM attendance WHERE DATE(timestamp) = ?",
        (target_date,),
    ).fetchall()
    present_uids = set(r[0] for r in present_rows)

    total_staff = len(all_users)
    present_count = len(present_uids)
    absent_count = max(0, total_staff - present_count)
    absent_users = [dict(u) for u in all_users if u["user_id"] not in present_uids]

    analytics = {
        "target_date": target_date,
        "total_staff": total_staff,
        "present_count": present_count,
        "absent_count": absent_count,
        "absent_users": absent_users,
    }

    # Table filtering
    conditions = []
    params = []

    if selected_user_id:
        conditions.append("a.user_id = ?")
        params.append(selected_user_id)

    if selected_from_date:
        conditions.append("DATE(a.timestamp) >= ?")
        params.append(selected_from_date)

    if selected_to_date:
        conditions.append("DATE(a.timestamp) <= ?")
        params.append(selected_to_date)

    where_clause = f"WHERE {' AND '.join(conditions)}" if conditions else ""

    count_sql = f"SELECT COUNT(*) FROM (SELECT 1 FROM attendance a {where_clause} GROUP BY a.user_id, DATE(a.timestamp))"
    total = db.execute(count_sql, params).fetchone()[0]

    rows_sql = f"""
        SELECT
            DATE(a.timestamp) AS date,
            a.user_id,
            COALESCE(u.name, 'Unknown') AS name,
            MIN(a.timestamp) AS min_ts,
            MAX(a.timestamp) AS max_ts,
            strftime('%I:%M:%S %p', MIN(a.timestamp)) AS in_time,
            strftime('%I:%M:%S %p', MAX(a.timestamp)) AS out_time,
            COUNT(*) AS total_punches,
            GROUP_CONCAT(strftime('%I:%M:%S %p', a.timestamp), ', ') AS all_punches
        FROM attendance a
        LEFT JOIN users u
            ON u.user_id = a.user_id
        {where_clause}
        GROUP BY a.user_id, DATE(a.timestamp)
        ORDER BY date DESC, CAST(a.user_id AS INTEGER), a.user_id
        LIMIT ? OFFSET ?
    """
    raw_rows = db.execute(rows_sql, params + [per_page, offset]).fetchall()

    db.close()

    rows = []
    for r in raw_rows:
        row_dict = dict(r)
        row_dict["total_hours"] = compute_total_hours(r["min_ts"], r["max_ts"], r["total_punches"])
        rows.append(row_dict)

    total_pages = max(1, (total + per_page - 1) // per_page)

    return render_template(
        "attendance.html",
        rows=rows,
        users_list=all_users,
        selected_user_id=selected_user_id,
        selected_from_date=selected_from_date,
        selected_to_date=selected_to_date,
        analytics=analytics,
        page=page,
        per_page=per_page,
        total=total,
        total_pages=total_pages,
    )


@app.route("/export")
def export_excel():
    selected_user_id = request.args.get("user_id", "").strip()
    selected_from_date = request.args.get("from_date", "").strip()
    selected_to_date = request.args.get("to_date", "").strip()
    preset = request.args.get("preset", "").strip()

    db = get_db()

    # Determine date range based on preset or default (latest month data)
    if preset:
        max_date_str = db.execute("SELECT MAX(DATE(timestamp)) FROM attendance").fetchone()[0] or datetime.now().strftime("%Y-%m-%d")
        max_dt = datetime.strptime(max_date_str, "%Y-%m-%d").date()

        if preset == "last_week":
            selected_from_date = (max_dt - timedelta(days=6)).strftime("%Y-%m-%d")
            selected_to_date = max_dt.strftime("%Y-%m-%d")
        elif preset == "last_month":
            first_of_curr = max_dt.replace(day=1)
            last_of_prev = first_of_curr - timedelta(days=1)
            first_of_prev = last_of_prev.replace(day=1)
            selected_from_date = first_of_prev.strftime("%Y-%m-%d")
            selected_to_date = last_of_prev.strftime("%Y-%m-%d")
    elif not selected_user_id and not selected_from_date and not selected_to_date:
        # Default: latest month data if no filter is active
        max_date_str = db.execute("SELECT MAX(DATE(timestamp)) FROM attendance").fetchone()[0]
        if max_date_str:
            max_dt = datetime.strptime(max_date_str, "%Y-%m-%d").date()
            selected_from_date = max_dt.replace(day=1).strftime("%Y-%m-%d")
            _, last_day = calendar.monthrange(max_dt.year, max_dt.month)
            selected_to_date = max_dt.replace(day=last_day).strftime("%Y-%m-%d")

    conditions = []
    params = []

    if selected_user_id:
        conditions.append("a.user_id = ?")
        params.append(selected_user_id)

    if selected_from_date:
        conditions.append("DATE(a.timestamp) >= ?")
        params.append(selected_from_date)

    if selected_to_date:
        conditions.append("DATE(a.timestamp) <= ?")
        params.append(selected_to_date)

    where_clause = f"WHERE {' AND '.join(conditions)}" if conditions else ""

    rows = db.execute(
        f"""
        SELECT
            DATE(a.timestamp) AS date,
            a.user_id,
            COALESCE(u.name, 'Unknown') AS name,
            MIN(a.timestamp) AS min_ts,
            MAX(a.timestamp) AS max_ts,
            strftime('%I:%M:%S %p', MIN(a.timestamp)) AS in_time,
            strftime('%I:%M:%S %p', MAX(a.timestamp)) AS out_time,
            COUNT(*) AS total_punches,
            GROUP_CONCAT(strftime('%I:%M:%S %p', a.timestamp), ', ') AS all_punches
        FROM attendance a
        LEFT JOIN users u
            ON u.user_id = a.user_id
        {where_clause}
        GROUP BY a.user_id, DATE(a.timestamp)
        ORDER BY date ASC, CAST(a.user_id AS INTEGER), a.user_id
        """,
        params,
    ).fetchall()

    db.close()

    # Create Excel Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Summary"

    # Styling definitions
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    headers = ["Date", "Name", "In Time (First)", "Out Time (Last)", "Total Hours", "Total Punches", "All Punches"]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    for row in rows:
        in_time_str = row["in_time"] if row["in_time"] else "-"
        out_time_str = row["out_time"] if row["total_punches"] > 1 else "- (Single Punch)"
        total_hours_str = compute_total_hours(row["min_ts"], row["max_ts"], row["total_punches"])
        ws.append([
            row["date"],
            row["name"],
            in_time_str,
            out_time_str,
            total_hours_str,
            row["total_punches"],
            row["all_punches"],
        ])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for i, cell in enumerate(row):
            cell.font = data_font
            cell.border = thin_border
            if i in [0, 2, 3, 4, 5]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    if preset:
        filename = f"attendance_{preset}_{selected_from_date}_to_{selected_to_date}.xlsx"
    elif selected_user_id and (selected_from_date or selected_to_date):
        filename = f"attendance_user_{selected_user_id}_{selected_from_date or 'start'}_to_{selected_to_date or 'end'}.xlsx"
    elif selected_user_id:
        filename = f"attendance_user_{selected_user_id}.xlsx"
    elif selected_from_date or selected_to_date:
        filename = f"attendance_{selected_from_date or 'start'}_to_{selected_to_date or 'end'}.xlsx"
    else:
        filename = f"attendance_latest_month_{selected_from_date}_to_{selected_to_date}.xlsx" if (selected_from_date and selected_to_date) else "attendance_latest_month.xlsx"

    return Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@app.route("/users")
def users():
    try:
        page = max(1, int(request.args.get("page", 1)))
    except ValueError:
        page = 1

    per_page = 25
    offset = (page - 1) * per_page

    db = get_db()

    total = db.execute(
        "SELECT COUNT(*) FROM users"
    ).fetchone()[0]

    rows = db.execute(
        """
        SELECT
            user_id,
            name,
            privilege,
            created_at,
            updated_at
        FROM users
        ORDER BY CAST(user_id AS INTEGER), user_id
        LIMIT ? OFFSET ?
        """,
        (per_page, offset),
    ).fetchall()

    db.close()

    total_pages = max(1, (total + per_page - 1) // per_page)

    return render_template(
        "users.html",
        rows=rows,
        page=page,
        total=total,
        total_pages=total_pages,
    )


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)

